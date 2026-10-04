"""Экспорт журнала инцидентов и сводной статистики в XLSX."""

import openpyxl
from openpyxl.styles import Font, PatternFill

from core.paths import data_path
from datetime import datetime


HEADER_FILL = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _style_header(ws, row=1):

    for cell in ws[row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT


def export_incidents_xlsx(logs):

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Инциденты"

    headers = ["Дата", "Источник", "Риск", "Действие", "AI-сервис", "Совпадений", "Типы данных", "Пользователь", "Файл"]
    ws.append(headers)
    _style_header(ws)

    for log in logs:
        ws.append([
            log.get("timestamp", ""),
            log.get("source", ""),
            log.get("risk", ""),
            log.get("action", "") or "",
            log.get("destination", "") or "",
            log.get("findings_count", 0),
            ", ".join(log.get("types", [])),
            log.get("created_by", "") or "",
            log.get("filename", "") or ""
        ])

    for i, width in enumerate([20, 12, 14, 22, 14, 12, 30, 16, 30], start=1):
        ws.column_dimensions[chr(64 + i)].width = width

    filename = datetime.now().strftime("PromptShield_Incidents_%Y%m%d_%H%M%S.xlsx")
    filepath = data_path("reports", filename)
    wb.save(filepath)

    return filepath


def export_summary_xlsx(stats, activity, period_label):

    wb = openpyxl.Workbook()

    ws1 = wb.active
    ws1.title = "Сводка"
    ws1.append(["Показатель", "Значение"])
    _style_header(ws1)
    ws1.append(["Период", period_label])
    ws1.append(["Проверок за период", stats.get("checks_period", 0)])
    ws1.append(["Проверок за всё время", stats.get("total_checks", 0)])
    ws1.append(["Доля высокого/критического риска, %", stats.get("high_risk_ratio", 0)])
    ws1.column_dimensions["A"].width = 35
    ws1.column_dimensions["B"].width = 20

    ws2 = wb.create_sheet("По уровню риска")
    ws2.append(["Уровень риска", "Количество"])
    _style_header(ws2)
    for risk, count in stats.get("risk_counts", {}).items():
        ws2.append([risk, count])
    ws2.column_dimensions["A"].width = 20
    ws2.column_dimensions["B"].width = 15

    ws3 = wb.create_sheet("Типы данных")
    ws3.append(["Тип", "Количество"])
    _style_header(ws3)
    for t, c in stats.get("top_types", []):
        ws3.append([t, c])
    ws3.column_dimensions["A"].width = 25
    ws3.column_dimensions["B"].width = 15

    ws4 = wb.create_sheet("Активность пользователей")
    ws4.append(["Пользователь", "Проверок", "Нарушений", "Оценка риска"])
    _style_header(ws4)
    for entry in activity:
        ws4.append([entry["username"], entry["total_checks"], entry["violations"], entry["risk_score"]])
    for col, width in zip("ABCD", [20, 12, 12, 14]):
        ws4.column_dimensions[col].width = width

    filename = datetime.now().strftime("PromptShield_Summary_%Y%m%d_%H%M%S.xlsx")
    filepath = data_path("reports", filename)
    wb.save(filepath)

    return filepath
